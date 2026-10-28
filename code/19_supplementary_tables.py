import json
import os

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

os.makedirs("supplementary", exist_ok=True)
N = json.load(open("results/numbers.json"))
NN = json.load(open("results/numbers_new.json"))
TITLE = "Interface removal reshapes the predicted homodimers of human splice isoforms while added sequence seldom builds interfaces"

pw = pd.read_csv("results/pairs_with_hc.tsv", sep="\t")
pa = pd.read_csv("results/pairs_analysis.tsv", sep="\t")
extra = ["canon_pred", "iso_pred", "canon_model_len", "iso_model_len", "model_len_ok", "both_pred", "canon_ipsae", "iso_ipsae",
         "canon_pdockq2", "iso_pdockq2", "canon_iptm", "iso_iptm", "canon_mono_plddt", "iso_mono_plddt", "grp"]
s1 = pw.merge(pa[["canonical_acc", "isoform_acc"] + extra], on=["canonical_acc", "isoform_acc"], how="left")
s1 = s1.drop(columns=["in_afdb_length_range"])
s1 = s1.rename(columns={"class": "class", "grp": "group", "ref_len": "canonical_length", "iso_len": "isoform_length",
                        "novel": "added_residues", "deleted": "removed_residues", "novel_N": "added_N", "novel_C": "added_C",
                        "novel_I": "added_internal", "deleted_N": "removed_N", "deleted_C": "removed_C",
                        "deleted_I": "removed_internal", "novel_segments": "added_segments", "deleted_segments": "removed_segments",
                        "canon_hc": "canonical_high_confidence", "iso_hc": "isoform_high_confidence",
                        "canon_pred": "canonical_predicted", "iso_pred": "isoform_predicted",
                        "canon_model_len": "canonical_model_chain_length", "iso_model_len": "isoform_model_chain_length",
                        "model_len_ok": "model_lengths_match_uniprot", "both_pred": "in_analysis_set",
                        "canon_ipsae": "canonical_ipSAE", "iso_ipsae": "isoform_ipSAE", "canon_pdockq2": "canonical_pDockQ2",
                        "iso_pdockq2": "isoform_pDockQ2", "canon_iptm": "canonical_ipTM", "iso_iptm": "isoform_ipTM",
                        "canon_mono_plddt": "canonical_monomer_pLDDT", "iso_mono_plddt": "isoform_monomer_pLDDT"})

rows = []
for name, key in [("Deposition call, analysis set", "transitions"),
                  ("Served models, all classified pairs", "transitions_all_pairs")]:
    for g, d in N[key].items():
        rows.append({"analysis": name, "group": g, "pairs": d["pairs"], "canonical_high_confidence": d["canon_hc"],
                     "losses": d["losses"], "loss_rate_pct": d["loss_rate"], "loss_ci_low": d["loss_ci"][0], "loss_ci_high": d["loss_ci"][1],
                     "canonical_not_high_confidence": d["canon_not_hc"], "gains": d["gains"], "gain_rate_pct": d["gain_rate"],
                     "gain_ci_low": d["gain_ci"][0], "gain_ci_high": d["gain_ci"][1]})
for g, d in N["transitions_ipsae_only"].items():
    rows.append({"analysis": "ipSAE threshold of 0.6 alone, analysis set", "group": g, "pairs": d["canon_below"] + d["canon_above"],
                 "canonical_high_confidence": d["canon_above"], "losses": d["losses"], "loss_rate_pct": d["loss_rate"],
                 "loss_ci_low": d["loss_ci"][0], "loss_ci_high": d["loss_ci"][1], "canonical_not_high_confidence": d["canon_below"],
                 "gains": d["gains"], "gain_rate_pct": d["gain_rate"], "gain_ci_low": d["gain_ci"][0], "gain_ci_high": d["gain_ci"][1]})
s3 = pd.DataFrame(rows).replace({"group": {"subtraction": "removal only", "addition": "addition only"}})

r1 = pd.read_csv("results/table_regression.tsv", sep="\t")
r1["set"] = "class-level models, analysis set"
r2 = pd.read_csv("results/table_retention_models.tsv", sep="\t")
r2["set"] = "interface retention models, analysis set with a high-confidence canonical"
s4 = pd.concat([r1, r2], ignore_index=True)[["set", "model", "n", "term", "OR", "lo", "hi", "p"]]
s4 = s4.rename(columns={"OR": "odds_ratio", "lo": "ci_low", "hi": "ci_high"})
s4[["odds_ratio", "ci_low", "ci_high"]] = s4[["odds_ratio", "ci_low", "ci_high"]].round(4)
s4["term"] = s4.term.str.replace("C(grp)[T.subtraction]", "removal only vs addition only", regex=False) \
    .str.replace("C(grp)[T.replacement]", "replacement vs addition only", regex=False)

s5 = pd.read_csv("results/canonical_retention.tsv", sep="\t")
s5 = s5.merge(pa[["canonical_acc", "isoform_acc", "iso_pred", "both_pred", "canon_ipsae", "iso_ipsae", "canon_pdockq2", "iso_pdockq2"]],
              on=["canonical_acc", "isoform_acc"], how="left").rename(columns={"both_pred": "in_analysis_set"})
s5 = s5.replace({"grp": {"subtraction": "removal only", "addition": "addition only"}}).rename(columns={"grp": "group"})

s6 = pd.read_csv("results/pair_mechanism.tsv", sep="\t").rename(columns={
    "novel_model": "added_residues_in_model", "novel_conf": "added_residues_pLDDT_ge70",
    "n_iface": "interface_residues", "n_iface_novel": "added_interface_residues",
    "n_iface_novel_conf": "added_interface_residues_pLDDT_ge70", "n_surface": "exposed_residues",
    "n_surface_novel": "exposed_added_residues", "novel_novel_contacts": "added_added_contacts",
    "mean_plddt_novel": "mean_pLDDT_added", "mean_plddt_iface_novel": "mean_pLDDT_added_interface",
    "canon_hc": "canonical_high_confidence"})

s7 = pd.read_csv("results/identical_copies.tsv", sep="\t")

gmap = {"addition": "addition only", "subtraction": "removal only", "replacement": "replacement", "all": "all"}
callname = {"deposition": "deposition (ipSAE>=0.6 and pDockQ2>=0.23)", "ipsae_only": "ipSAE>=0.6 alone",
            "pdockq2_only": "pDockQ2>=0.23 alone", "iptm_ge08": "ipTM>=0.8", "iptm_ge06": "ipTM>=0.6"}
rob_rows = []
for call, d in NN["robustness"].items():
    if call == "iptm_vs_call":
        continue
    label = callname.get(call, call.replace("grid_ipsae", "ipSAE>=").replace("_pdockq2", ", pDockQ2>="))
    for g in ["addition", "subtraction", "replacement", "all"]:
        v = d[g]
        rob_rows.append({"call": label, "group": gmap[g], "canonical_high_confidence": v["canon_hc"],
                         "losses": v["losses"], "loss_rate_pct": v["loss_rate"],
                         "canonical_not_high_confidence": v["canon_not"], "gains": v["gains"], "gain_rate_pct": v["gain_rate"]})
s_rob = pd.DataFrame(rob_rows)

s_swap = pd.read_csv("results/swap_segments.tsv", sep="\t").rename(columns={
    "site": "swap_site", "can_start": "removed_start", "can_end": "removed_end", "iso_start": "added_start",
    "iso_end": "added_end", "len_removed": "removed_length", "len_added": "added_length", "score": "alignment_score",
    "aligned": "aligned_columns", "identity": "aligned_identity", "z": "z_score", "homologous": "homologous_swap",
    "removed_mean_plddt": "removed_segment_mean_pLDDT", "added_mean_plddt": "added_segment_mean_pLDDT"})

s_added = pd.read_csv("results/added_residues_monomer.tsv.gz", sep="\t").rename(columns={
    "cls": "group", "n_added": "added_residues", "added_lt50": "added_fraction_pLDDT_lt50",
    "added_ge70": "added_fraction_pLDDT_ge70", "added_mean": "added_mean_pLDDT", "n_shared": "shared_residues",
    "shared_lt50": "shared_fraction_pLDDT_lt50", "shared_ge70": "shared_fraction_pLDDT_ge70",
    "n_added_conf": "added_residues_pLDDT_ge70", "added_N": "added_N_terminal", "added_C": "added_C_terminal",
    "added_I": "added_internal"}).replace({"group": {"subtraction": "removal only", "addition": "addition only"}})
for c in ["added_fraction_pLDDT_lt50", "added_fraction_pLDDT_ge70", "added_mean_pLDDT", "shared_fraction_pLDDT_lt50", "shared_fraction_pLDDT_ge70"]:
    s_added[c] = s_added[c].round(4)

sheets = [
    ("Table S1", s1, f"All {N['pairs_in_range']:,} human isoform-canonical pairs with both sequences of 15 to 1,500 residues: class, added and removed residues, manifest scores (larger chain direction), model chain lengths and high-confidence calls. The analysis set ({N['pairs_analysed']:,} pairs) has both homodimers predicted with the chain lengths of the current UniProt sequences."),
    ("Table S2", s3, "Transitions of the high-confidence call by isoform class under the deposition call, under the ipSAE threshold alone, and from served models over all classified pairs; rates in percent with Wilson 95% confidence intervals."),
    ("Table S3", s4, "Logistic models for gain and loss of the high-confidence call; odds ratios with Wald 95% confidence intervals. Interface and residue terms are per 10 percentage points; d_plddt is per pLDDT point."),
    ("Table S4", s5, f"Interface retention for the {len(s5):,} pairs with a high-confidence canonical homodimer whose model sequence matches UniProt; the {N['retention_pairs']:,} in the analysis set enter the retention analysis."),
    ("Table S5", s6, f"The {len(s6):,} high-confidence isoform homodimers in the classified pairs with their interface composition."),
    ("Table S6", s7, f"The {len(s7):,} identical-sequence copies of {N['copies_human_with_copy']} sampled human canonical sequences, with deposition status, manifest values (larger chain direction) and, for copies served by the API, the API's two-decimal ipSAE and pDockQ2."),
    ("Table S7", s_rob, "Transitions of the high-confidence call by isoform class under alternative definitions of the call: the deposition rule, the ipSAE threshold alone, the pDockQ2 threshold alone, ipTM cutoffs of 0.8 and 0.6, and a grid of ipSAE and pDockQ2 thresholds. Rates in percent over the analysis set."),
    ("Table S8", s_swap, f"The {len(s_swap):,} replacement swaps (a gap between shared blocks where the isoform both removes and adds residues): the removed canonical segment, the added isoform segment, their local alignment (BLOSUM62) with z-score against shuffled controls, the homology call, and the mean monomer pLDDT of each segment."),
    ("Table S9", s_added, f"Added-sequence confidence in the monomer models for each of the {len(s_added):,} isoforms that carry added residues and have an AlphaFold monomer model: counts and pLDDT fractions of added and shared residues, and the location of the added residues."),
]
out = "supplementary/Supplementary_Tables.xlsx"
with pd.ExcelWriter(out, engine="openpyxl") as xw:
    tp = pd.DataFrame({"": ["Supplementary Information", "", "Title:", "Authors:", ""] + [s for s, _, _ in sheets],
                       " ": ["", "", TITLE, "Taner Karagöl, Alper Karagöl", ""] + [d for _, _, d in sheets]})
    tp.to_excel(xw, sheet_name="Title Page", index=False, header=False)
    for name, df, _ in sheets:
        df.to_excel(xw, sheet_name=name, index=False)
    for ws in xw.book.worksheets:
        if ws.title == "Title Page":
            ws.column_dimensions["A"].width = 26
            ws.column_dimensions["B"].width = 140
            ws["A1"].font = Font(bold=True)
            for r in ws.iter_rows(min_row=1, max_row=ws.max_row):
                for c in r:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            continue
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 48)
print("wrote", out, {name: len(df) for name, df, _ in sheets})
